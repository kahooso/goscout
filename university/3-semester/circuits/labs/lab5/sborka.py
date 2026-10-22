"""ЛР-5: сборка Word-отчёта «{{STUDENT_FILE}} Лаб5.docx».

Шаблон — отчёт ЛР-4: от него берутся титульный лист и стили (ОТЧЁТ_*).
Всё начиная с «Цель работы» собирается заново, поэтому правки руками
после пересборки пропадут — сначала вписывать числа в zamery.py и пути в SKRINY.

Пока числа нет — в отчёте жёлтое «___», пока скрина нет — жёлтая пометка
[СКРИН Sxx: …]. Графики строятся из zamery.py (grafik_*.png).

Запуск: python sborka.py
"""
import importlib
import re
import shutil
import time
from pathlib import Path

import docx
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm

LAB = Path(r"C:/Users/ryche/Desktop/goscout/university/3-semester/circuits/labs/lab5")
SHABLON = LAB.parent / "lab4" / "{{STUDENT_FILE}} Лаб4.docx"
DOCX = LAB / "{{STUDENT_FILE}} Лаб5.docx"
SS = Path(r"D:/Pictures/Screenshots")

import sys
sys.path.insert(0, str(LAB))
Z = importlib.import_module("zamery")


def ss(stamp, day="2026-10-13"):
    return SS / f"Снимок экрана {day} {stamp}.png"


# код → (путь или None, ширина в см, что снять)
SKRINY = {
    "S01": (None, 16, "схема стенда (рис. 5 методички), S2/S3 «прямое»"),
    "S02": (ss("221905"), 16, "схема источника питания (рис. 6 методички)"),
    "S03": (ss("221232"), 16, "стенд: D1 прямое, R1 = 50 %"),
    "S04": (None, 16, "стенд: D2 прямое, R1 = 50 %, окна XMM1 и XMM2 открыты"),
    "S05": (None, 16, "стенд: D3 прямое, R1 = 50 %, окна XMM1 и XMM2 открыты"),
    "S06": (None, 16, "стенд: D1 обратное, V1 = 50 В, R1 = 90 %"),
    "S07": (None, 16, "стенд: D2 обратное, V1 = 50 В, R1 = 90 %"),
    "S08": (None, 16, "стенд: D3 обратное, V1 = 15 В, R1 = 20 % (пробой)"),
    "S09": (None, 16, "осциллограф: без C1 (конденсатор удалён), 50 Гц"),
    "S10": (None, 16, "осциллограф: C1 = 300 мкФ, 50 Гц"),
    "S11": (None, 16, "осциллограф: C1 = 30 мкФ, 50 Гц"),
    "S12": (None, 16, "осциллограф: C1 = 30 мкФ, 500 Гц"),
    "S13": (None, 16, "схема источника со стабилитроном (Rд = 1 кОм + D 02BZ2.2)"),
    "S14": (None, 16, "осциллограф: источник со стабилитроном, C1 = 300 мкФ, 50 Гц"),
}

# ---------------------------------------------------------------------------
# числа
# ---------------------------------------------------------------------------
NONE = "___"


def f(x, nd=3):
    if x is None:
        return NONE
    return f"{x:.{nd}f}".replace(".", ",")


def g(x, sig=3):
    """3 значащие цифры, запятая"""
    if x is None:
        return NONE
    if x == 0:
        return "0"
    if abs(x) >= 1e5:
        e = len(str(int(abs(x)))) - 1
        sup = str(e).translate(str.maketrans("0123456789-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻"))
        return f"{x / 10 ** e:.{sig - 1}f}".replace(".", ",") + f"·10{sup}"
    s = f"{x:.{sig}g}"
    if "e" in s:
        s = f"{x:.{max(0, sig - 1)}f}"
    return s.replace(".", ",")


def om(r):
    """сопротивление с приставкой"""
    for k, name in ((1e9, "ГОм"), (1e6, "МОм"), (1e3, "кОм")):
        if abs(r) >= k:
            return f"{g(r / k)} {name}"
    return f"{g(r)} Ом"


def unit_for(rows):
    vals = [abs(r[2]) for r in rows if r[2] is not None]
    m = max(vals) if vals else 1e-3
    if m >= 1e-3:
        return "мА", 1e3
    if m >= 1e-6:
        return "мкА", 1e6
    return "нА", 1e9


def gotovo(rows):
    return [r for r in rows if r[1] is not None and r[2] is not None]


def interp_u(rows, i_target):
    """U при заданном токе — линейная интерполяция между соседними точками"""
    pts = sorted(gotovo(rows), key=lambda r: r[2])
    for a, b in zip(pts, pts[1:]):
        if a[2] <= i_target <= b[2] and b[2] != a[2]:
            return a[1] + (b[1] - a[1]) * (i_target - a[2]) / (b[2] - a[2])
    return None


def rab_tochka(rows, i_target):
    """рабочая точка: точка с током ближе всего к i_target, у которой есть соседи.
    Возвращает (U, I, Rст, rдиф, U-, U+, I-, I+) или None"""
    pts = gotovo(rows)
    if len(pts) < 3:
        return None
    k = min(range(1, len(pts) - 1), key=lambda j: abs(pts[j][2] - i_target))
    a, p, b = pts[k - 1], pts[k], pts[k + 1]
    if p[2] == 0 or b[2] == a[2]:
        return None
    return (p[1], p[2], p[1] / p[2], (b[1] - a[1]) / (b[2] - a[2]), a[1], b[1], a[2], b[2])


DIODY = {"D1": "1BH62", "D2": "1N1199C", "D3": "02BZ2.2"}

RT = {d: rab_tochka(Z.PRYAM[d], 20e-3) for d in DIODY}
U_OTKR = {d: interp_u(Z.PRYAM[d], 1e-3) for d in DIODY}
RT_ST = rab_tochka(Z.STAB, 5e-3)
U_ST = interp_u(Z.STAB, 5e-3)


def obr_max(d):
    pts = gotovo(Z.OBR[d])
    if not pts:
        return None
    return max(pts, key=lambda r: r[1])


def st_diapazon():
    """первая и последняя точки на участке пробоя (I ≥ 1 мА)"""
    pts = [r for r in gotovo(Z.STAB) if r[2] >= 1e-3]
    if len(pts) < 2:
        return None
    return pts[0], pts[-1]


def istochnik_izm():
    rows = Z.STAB_ISTOCHNIK
    keys = sorted(rows)
    a, b = rows[keys[0]], rows[keys[-1]]
    if None in a or None in b:
        return None
    return keys[0], keys[-1], b[0] - a[0], b[1] - a[1]


# ---------------------------------------------------------------------------
# графики
# ---------------------------------------------------------------------------
CVET = {"D1": "#1f77b4", "D2": "#d62728", "D3": "#2ca02c"}


def grafik_pryam():
    path = LAB / "grafik_1_pryamye.png"
    if not any(gotovo(Z.PRYAM[d]) for d in DIODY):
        return None
    fig, ax = plt.subplots(figsize=(7, 4.5), dpi=150)
    for d, name in DIODY.items():
        pts = gotovo(Z.PRYAM[d])
        if pts:
            ax.plot([r[1] for r in pts], [r[2] * 1e3 for r in pts], "o-",
                    color=CVET[d], ms=4, label=f"{d} ({name})")
    ax.set_xlabel("U, В")
    ax.set_ylabel("I, мА")
    ax.grid(True, alpha=0.4)
    ax.legend()
    ax.set_xlim(left=0)
    ax.set_ylim(bottom=0)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


def grafik_obr():
    path = LAB / "grafik_2_obratnye.png"
    allpts = [r for d in Z.OBR for r in gotovo(Z.OBR[d])]
    if not allpts:
        return None
    uname, k = unit_for(allpts)
    fig, ax = plt.subplots(figsize=(7, 4.5), dpi=150)
    for d in Z.OBR:
        pts = gotovo(Z.OBR[d])
        if pts:
            ax.plot([-r[1] for r in pts], [-r[2] * k for r in pts], "o-",
                    color=CVET[d], ms=4, label=f"{d} ({DIODY[d]})")
    ax.set_xlabel("U, В")
    ax.set_ylabel(f"I, {uname}")
    ax.grid(True, alpha=0.4)
    ax.legend()
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


def grafik_stab():
    path = LAB / "grafik_3_stabilitron.png"
    pts = gotovo(Z.STAB)
    if not pts:
        return None
    fig, ax = plt.subplots(figsize=(7, 4.5), dpi=150)
    ax.plot([-r[1] for r in pts], [-r[2] * 1e3 for r in pts], "o-",
            color=CVET["D3"], ms=4, label="D3 (02BZ2.2), обратное включение")
    if U_ST is not None:
        ax.axvline(-U_ST, color="gray", ls="--", lw=1)
        ax.text(-U_ST, ax.get_ylim()[0] * 0.5, f"  Uст ≈ {f(U_ST, 2)} В", color="gray")
    ax.set_xlabel("U, В")
    ax.set_ylabel("I, мА")
    ax.grid(True, alpha=0.4)
    ax.legend(loc="upper left")
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


# ---------------------------------------------------------------------------
# таблицы
# ---------------------------------------------------------------------------

def tabl_pryam():
    head = ["R1, %"] + [x for d in DIODY for x in (f"U{d[1]}, В", f"I{d[1]}, мА")]
    rows = [head]
    for i, p in enumerate(Z.POZ_PRYAM):
        row = [str(p)]
        for d in DIODY:
            _, u, cur = Z.PRYAM[d][i]
            row += [f(u, 3), g(None if cur is None else cur * 1e3)]
        rows.append(row)
    return rows


def tabl_obr():
    allpts = [r for d in Z.OBR for r in gotovo(Z.OBR[d])]
    uname, k = unit_for(allpts) if allpts else ("мкА", 1e6)
    rows = [["R1, %", "U1, В", f"I1, {uname}", "U2, В", f"I2, {uname}"]]
    for i, p in enumerate(Z.POZ_OBR):
        row = [str(p)]
        for d in ("D1", "D2"):
            _, u, cur = Z.OBR[d][i]
            row += [f(u, 2), g(None if cur is None else cur * k)]
        rows.append(row)
    return rows


def tabl_stab():
    rows = [["R1, %", "U, В", "I, мА"]]
    for p, u, cur in Z.STAB:
        rows.append([str(p), f(u, 3), g(None if cur is None else cur * 1e3)])
    return rows


def tabl_sopr():
    rows = [["Диод", "U, В", "I, мА", "Rст = U/I, Ом", "rдиф = ΔU/ΔI, Ом"]]
    for d, name in DIODY.items():
        t = RT[d]
        if t is None:
            rows.append([f"{d} ({name}), прямое", NONE, NONE, NONE, NONE])
        else:
            rows.append([f"{d} ({name}), прямое", f(t[0], 3), g(t[1] * 1e3),
                         g(t[2]), g(t[3])])
    t = RT_ST
    if t is None:
        rows.append(["D3 (02BZ2.2), пробой", NONE, NONE, NONE, NONE])
    else:
        rows.append(["D3 (02BZ2.2), пробой", f(t[0], 3), g(t[1] * 1e3), g(t[2]), g(t[3])])
    return rows


def tabl_istochnik():
    rows = [["Амплитуда генератора, В", "Uвых без стабилитрона, В", "Uвых со стабилитроном, В"]]
    for a in sorted(Z.STAB_ISTOCHNIK):
        b, c = Z.STAB_ISTOCHNIK[a]
        rows.append([str(a), f(b, 2), f(c, 3)])
    return rows


# ---------------------------------------------------------------------------
# тексты
# ---------------------------------------------------------------------------

def txt_pryam_itog():
    u = {d: f(U_OTKR[d], 2) for d in DIODY}
    return ("На всех трёх графиках при малых напряжениях ток почти не течёт. После "
            "определённого напряжения ток начинает быстро расти, а дальше рост становится "
            "почти линейным. Напряжение, при котором ток достигает 1 мА: "
            f"D1 — {u['D1']} В, D2 — {u['D2']} В, D3 — {u['D3']} В. До этого "
            "напряжения диод практически закрыт: внешнее напряжение ещё не преодолело "
            "контактную разность потенциалов p-n-перехода.")


def txt_sopr():
    r = {d: (NONE if RT[d] is None else g(RT[d][2])) for d in DIODY}
    rd = {d: (NONE if RT[d] is None else g(RT[d][3])) for d in DIODY}
    return ("Статическое сопротивление в рабочей точке: D1 — "
            f"{r['D1']} Ом, D2 — {r['D2']} Ом, D3 — {r['D3']} Ом; динамическое: "
            f"D1 — {rd['D1']} Ом, D2 — {rd['D2']} Ом, D3 — {rd['D3']} Ом. "
            "Динамическое сопротивление меньше статического: на рабочем участке ВАХ "
            "небольшое изменение напряжения вызывает большое изменение тока. "
            "Сопротивление диода не постоянно — оно уменьшается с ростом тока, "
            "то есть диод является нелинейным элементом.")


def txt_obr():
    parts = []
    for d in ("D1", "D2"):
        m = obr_max(d)
        rp = RT[d]
        if m is None or m[2] == 0:
            parts.append(f"{d}: при U = {NONE} В обратный ток {NONE}, "
                         f"обратное сопротивление {NONE}")
            continue
        uname, k = unit_for([m])
        r_obr = m[1] / m[2]
        s = (f"{d}: при U = {f(m[1], 1)} В обратный ток {g(m[2] * k)} {uname}, "
             f"обратное сопротивление U/I ≈ {om(r_obr)}")
        if rp is not None:
            s += f" — примерно в {g(r_obr / rp[2], 2)} раз больше прямого"
        parts.append(s)
    return ("; ".join(parts) + ". Обратный ток во много раз меньше прямого: "
            "при обратном включении потенциальный барьер p-n-перехода растёт, "
            "и через переход идёт только малый ток неосновных носителей. "
            "Пробой: " + NONE + ".")


def txt_stab():
    dia = st_diapazon()
    if dia is None:
        izm = (f"При изменении тока от {NONE} до {NONE} мА напряжение на стабилитроне "
               f"изменилось только с {NONE} до {NONE} В.")
    else:
        a, b = dia
        izm = (f"При изменении тока от {g(a[2] * 1e3)} до {g(b[2] * 1e3)} мА напряжение "
               f"на стабилитроне изменилось только с {f(a[1], 3)} до {f(b[1], 3)} В.")
    rd = NONE if RT_ST is None else g(RT_ST[3])
    return ("До напряжения пробоя ток стабилитрона мал. После пробоя ток резко растёт, "
            "а напряжение почти не меняется: " + izm + " Напряжение стабилизации "
            f"(при токе 5 мА) Uст ≈ {f(U_ST, 2)} В, что соответствует номиналу "
            f"стабилитрона 2,2 В. Динамическое сопротивление на участке пробоя "
            f"rдиф ≈ {rd} Ом — оно мало, поэтому напряжение и держится постоянным.")


def txt_istochnik_izm():
    r = istochnik_izm()
    if r is None:
        return (f"При увеличении амплитуды входного сигнала с {NONE} до {NONE} В "
                f"выходное напряжение без стабилитрона выросло на {NONE} В, "
                f"а со стабилитроном — только на {NONE} В.")
    a, b, d1, d2 = r
    return (f"При увеличении амплитуды входного сигнала с {a} до {b} В выходное "
            f"напряжение без стабилитрона выросло на {f(d1, 2)} В, а со стабилитроном — "
            f"только на {f(d2, 3)} В.")


def content():
    c = []
    c += [
        ("h2", "Исследование прямых ветвей ВАХ диодов"),
        ("p", "Схема стенда приведена на рисунке 2.1. Переключатели S2 и S3 установлены "
              "в положение «прямое»: плюс источника подан на анод диода. Напряжение "
              "источника V1 = 10 В. Исследуемый диод выбирается ключами S4–S6 (замкнут "
              "только один). Напряжение на диоде меняется потенциометром R1 от 0 до 80 %: "
              "при 100 % диод подключается к источнику напрямую и ток ничем не ограничен."),
        ("p", "Пример измерения для каждого диода (R1 = 50 %) показан на рисунках "
              "{s_d1}–{s_d3}."),
        ("fig", "s_d1", "S03", "Прямое включение диода D1 (1BH62)"),
        ("fig", "s_d2", "S04", "Прямое включение диода D2 (1N1199C)"),
        ("fig", "s_d3", "S05", "Прямое включение стабилитрона D3 (02BZ2.2)"),
        ("p", "Результаты измерений приведены в таблице {t_pr} (индекс 1, 2, 3 — номер "
              "диода), ВАХ — на рисунке {g_pr}."),
        ("tbl", "t_pr", "Прямые ветви ВАХ диодов D1, D2, D3", tabl_pryam()),
        ("fig", "g_pr", "GPR", "Прямые ветви ВАХ диодов"),
        ("p", txt_pryam_itog()),

        ("h2", "Расчёт статического и динамического сопротивлений"),
        ("p", "Рабочая точка для каждого диода выбрана на рабочем (крутом) участке ВАХ — "
              "при токе около 20 мА. Статическое сопротивление считается по формуле 1, "
              "динамическое — по формуле 2, где ΔU и ΔI взяты между соседними с рабочей "
              "точкой отсчётами таблицы {t_pr}."),
    ]
    t = RT["D1"]
    if t is None:
        c += [("eq", "Rст = U/I = ___ / ___ = ___ Ом"),
              ("eq", "rдиф = ΔU/ΔI = (___ − ___) / (___ − ___) = ___ Ом")]
    else:
        c += [("p", "Пример расчёта для диода D1:"),
              ("eq", f"Rст = U/I = {f(t[0], 3)} В / {g(t[1] * 1e3)} мА = {g(t[2])} Ом"),
              ("eq", f"rдиф = ΔU/ΔI = ({f(t[5], 3)} − {f(t[4], 3)}) В / "
                     f"({g(t[7] * 1e3)} − {g(t[6] * 1e3)}) мА = {g(t[3])} Ом")]
    c += [
        ("p", "Результаты для всех диодов приведены в таблице {t_sopr}."),
        ("tbl", "t_sopr", "Статические и динамические сопротивления в рабочих точках",
         tabl_sopr()),
        ("p", txt_sopr()),

        ("h2", "Исследование обратных ветвей ВАХ диодов D1 и D2"),
        ("p", "Переключатели S2 и S3 переведены в положение «обратное»: плюс источника "
              f"подан на катод. Напряжение источника V1 = {Z.V_OBR} В. Примеры измерения "
              "показаны на рисунках {s_o1} и {s_o2}, результаты — в таблице {t_obr}, "
              "ВАХ — на рисунке {g_obr} (обратные напряжение и ток отложены со знаком минус)."),
        ("fig", "s_o1", "S06", "Обратное включение диода D1"),
        ("fig", "s_o2", "S07", "Обратное включение диода D2"),
        ("tbl", "t_obr", "Обратные ветви ВАХ диодов D1 и D2", tabl_obr()),
        ("fig", "g_obr", "GOBR", "Обратные ветви ВАХ диодов D1 и D2"),
        ("p", txt_obr()),

        ("h2", "Исследование обратной ветви ВАХ стабилитрона"),
        ("p", "Ключом S6 выбран стабилитрон D3, переключатели S2 и S3 — в положении "
              "«обратное», V1 = 15 В. Вблизи пробоя (R1 = 10–20 %) потенциометр "
              "менялся с шагом 1–2 %. Пример измерения в режиме пробоя — рисунок "
              "{s_st}, результаты — таблица {t_st}, ВАХ — рисунок {g_st}."),
        ("fig", "s_st", "S08", "Стабилитрон D3 в режиме пробоя"),
        ("tbl", "t_st", "Обратная ветвь ВАХ стабилитрона D3", tabl_stab()),
        ("fig", "g_st", "GST", "Обратная ветвь ВАХ стабилитрона D3"),
        ("p", txt_stab()),

        ("h2", "Исследование источника питания"),
        ("p", "Схема источника питания приведена на рисунке 2.2. Функциональный "
              "генератор XFG1 выдаёт синусоидальное напряжение частотой 50 Гц; амплитуда "
              "генератора 5 В, между выводами «+» и «−» — 10 В. Трансформатор T1 передаёт его на диодный мост D1 (3N259), который "
              "выпрямляет обе полуволны. Конденсатор C1 вместе с нагрузкой R1 = 5 кОм "
              "образует фильтр нижних частот и сглаживает выпрямленное напряжение. "
              "Канал A осциллографа — входной сигнал, канал B — напряжение на нагрузке."),
        ("p", "Сначала посмотрим выпрямленное напряжение без конденсатора:"),
        ("fig", "o_bezc", "S09", "Напряжение на нагрузке без конденсатора"),
        ("p", "Без конденсатора на нагрузке пульсирующее напряжение одной полярности: "
              "обе полуволны синусоиды «перевёрнуты» вверх, напряжение падает до нуля "
              "100 раз в секунду (вдвое чаще частоты сигнала)."),
        ("p", "Подключим конденсатор 300 мкФ:"),
        ("fig", "o_c300", "S10", "Напряжение на нагрузке, C1 = 300 мкФ, 50 Гц"),
        ("p", f"Напряжение стало почти постоянным, размах пульсаций {NONE} В. "
              "Конденсатор заряжается на вершинах полуволн и между ними медленно "
              "разряжается через нагрузку."),
        ("p", "Уменьшим ёмкость до 30 мкФ:"),
        ("fig", "o_c30", "S11", "Напряжение на нагрузке, C1 = 30 мкФ, 50 Гц"),
        ("p", f"Размах пульсаций вырос до {NONE} В: меньший конденсатор быстрее "
              "разряжается через нагрузку (постоянная времени τ = R·C уменьшилась "
              "в 10 раз)."),
        ("p", "Увеличим частоту генератора до 500 Гц при той же ёмкости 30 мкФ:"),
        ("fig", "o_f500", "S12", "Напряжение на нагрузке, C1 = 30 мкФ, 500 Гц"),
        ("p", f"Размах пульсаций уменьшился до {NONE} В: промежуток между вершинами "
              "стал в 10 раз короче, и конденсатор не успевает разрядиться."),
        ("p", "Преобразуем схему в стабилизированный источник: между конденсатором "
              "и нагрузкой включим гасящий резистор Rд = 1 кОм, а параллельно нагрузке — "
              "стабилитрон 02BZ2.2 в обратном включении (рисунок {sh_st})."),
        ("fig", "sh_st", "S13", "Схема стабилизированного источника питания"),
        ("fig", "o_st", "S14", "Напряжение на нагрузке стабилизированного источника"),
        ("p", f"Напряжение на нагрузке стало равно {NONE} В — это напряжение "
              "стабилизации стабилитрона, пульсации практически исчезли. Избыток "
              "напряжения падает на резисторе Rд."),
        ("p", "Для оценки стабилизации амплитуда генератора менялась от 3 до 7 В, "
              "выходное напряжение измерялось мультиметром (таблица {t_ist})."),
        ("tbl", "t_ist", "Выходное напряжение источника при разной амплитуде входного сигнала",
         tabl_istochnik()),
        ("p", txt_istochnik_izm() + " Схема со стабилитроном удерживает выходное "
              "напряжение почти постоянным при изменении входного."),
    ]
    return c


def vyvody():
    u = {d: f(U_OTKR[d], 2) for d in DIODY}
    return [
        "Диод проводит ток только в одном направлении. В прямом включении ток "
        "практически отсутствует до напряжения открывания (ток 1 мА достигается при "
        f"{u['D1']} В у D1, {u['D2']} В у D2, {u['D3']} В у D3), после чего быстро "
        "растёт; дальше рост становится почти линейным из-за сопротивления самих "
        "p- и n-областей.",
        "В обратном включении ток диода очень мал, обратное сопротивление во много раз "
        "больше прямого.",
        "Сопротивление диода зависит от рабочей точки — диод является нелинейным "
        "элементом. Динамическое сопротивление на рабочем участке меньше статического.",
        f"Стабилитрон 02BZ2.2 в обратном включении пробивается при напряжении около "
        f"{f(U_ST, 2)} В. На участке пробоя ток меняется в широких пределах, а "
        "напряжение почти постоянно — это и используется для стабилизации.",
        "Диодный мост выпрямляет обе полуволны переменного напряжения. Конденсатор "
        "сглаживает пульсации: чем больше ёмкость и выше частота, тем пульсации меньше.",
        "Стабилитрон с гасящим резистором удерживает выходное напряжение источника "
        "почти постоянным при изменении входного напряжения. Недостатки такой схемы: "
        "выходное напряжение задаётся типом стабилитрона, на резисторе Rд теряется "
        "мощность, ток нагрузки ограничен.",
    ]


TEORIYA = [
    ("p", "Полупроводники — материалы, которые проводят ток лучше диэлектриков, но хуже "
          "проводников (кремний, германий, арсенид галлия). Добавляя примеси, получают "
          "полупроводник n-типа (основные носители заряда — электроны) или p-типа "
          "(основные носители — дырки)."),
    ("p", "На границе p- и n-областей образуется p-n-переход. Электроны и дырки "
          "диффундируют через границу и рекомбинируют, рядом с границей остаются "
          "неподвижные ионы примесей. Образуется обеднённый слой с электрическим полем, "
          "которое препятствует дальнейшей диффузии, — контактная разность потенциалов."),
    ("p", "Полупроводниковый диод — прибор с одним p-n-переходом и двумя выводами: анод "
          "(p-область) и катод (n-область). При прямом включении (плюс на аноде) внешнее "
          "напряжение уменьшает барьер, и ток быстро растёт. При обратном включении барьер "
          "растёт, через диод идёт только малый обратный ток. Зависимость тока от "
          "напряжения называется вольт-амперной характеристикой (ВАХ)."),
    ("p", "При большом обратном напряжении наступает пробой — резкий рост обратного тока. "
          "Различают лавинный, туннельный и тепловой пробой. Лавинный и туннельный пробой "
          "обратимы, тепловой разрушает переход."),
    ("p", "Стабилитрон — диод, рассчитанный на работу в режиме обратимого пробоя. На этом "
          "участке ВАХ напряжение почти не зависит от тока, поэтому стабилитрон "
          "используется для получения стабильного напряжения. Ток через него "
          "ограничивается последовательным резистором Rд."),
    ("p", "Так как ВАХ диода нелинейна, его сопротивление зависит от рабочей точки. "
          "Статическое сопротивление рассчитывается по формуле 1:"),
    ("eq", "Rст = U / I"),
    ("p", "где U — напряжение на диоде в рабочей точке, В; I — ток через диод, А. "
          "Динамическое (дифференциальное) сопротивление показывает, насколько "
          "меняется напряжение при небольшом изменении тока, формула 2:"),
    ("eq", "rдиф = ΔU / ΔI"),
    ("p", "где ΔU — изменение напряжения, В; ΔI — соответствующее изменение тока, А."),
    ("p", "Источник питания состоит из трансформатора, выпрямителя и фильтра. Диодный "
          "мост пропускает на нагрузку обе полуволны синусоиды в одном направлении. "
          "Конденсатор, подключённый параллельно нагрузке, заряжается на вершинах "
          "полуволн и разряжается через нагрузку между ними; чем больше постоянная "
          "времени τ = R·C по сравнению с периодом пульсаций, тем меньше пульсации."),
]

# ---------------------------------------------------------------------------
# сборка docx
# ---------------------------------------------------------------------------
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
GRAFIKI = {}
VSTAVLENO = {}   # подпись рисунка → файл картинки, вставленной руками в Word
SKR_DIR = LAB / "skriny"


def sobrat_vstavlennye():
    """Картинки, которые пользователь вставил в Word сам, сохраняются в skriny/
    под именем подписи — при пересборке они встают на то же место."""
    if not DOCX.exists():
        return
    SKR_DIR.mkdir(exist_ok=True)
    d = docx.Document(str(DOCX))
    ps = d.paragraphs
    for i, p in enumerate(ps[:-1]):
        if p.style.name != "ОТЧЁТ_Рисунок":
            continue
        blips = p._p.xpath(".//a:blip/@r:embed")
        cap = ps[i + 1].text.strip().lstrip("–").strip()
        if not blips or not cap:
            continue
        part = d.part.related_parts[blips[0]]
        ext = Path(part.partname).suffix
        name = re.sub(r'[\/:*?"<>|]', "_", cap)
        out = SKR_DIR / f"{name}{ext}"
        out.write_bytes(part.blob)
        VSTAVLENO[cap] = out


def add_text(doc, text, style="Normal"):
    p = doc.add_paragraph(style=style)
    for part in re.split(r"(___|\[СКРИН[^\]]*\]|\[ГРАФИК[^\]]*\])", text):
        if not part:
            continue
        r = p.add_run(part)
        if part == NONE or part.startswith("["):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            r.bold = True
    return p


def add_eq(doc, text):
    p = doc.add_paragraph(style="Normal")
    esc = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    p._p.append(parse_xml(
        f'<m:oMathPara xmlns:m="{M_NS}"><m:oMath><m:r>'
        f'<m:t xml:space="preserve">{esc}</m:t></m:r></m:oMath></m:oMathPara>'))


def add_fig(doc, src, caption, label):
    if src in GRAFIKI:
        path, width = GRAFIKI[src], 15
        hint = "ГРАФИК: появится после замеров"
    else:
        path, width, what = SKRINY[src]
        hint = f"СКРИН {label}: {what}"
        if caption in VSTAVLENO:
            path = VSTAVLENO[caption]
    p = doc.add_paragraph(style="ОТЧЁТ_Рисунок")
    if path is not None and Path(path).exists():
        p.add_run().add_picture(str(path), width=Cm(width))
    else:
        r = p.add_run(f"[{hint}]")
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        r.bold = True
    doc.add_paragraph("– " + caption, style="ОТЧЁТ_Рисунок_подпись")


def add_table(doc, caption, rows):
    cap = doc.add_paragraph(caption, style="ОТЧЁТ_Таблица_подпись")
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = doc.styles["ОТЧЁТ_Таблица"]
    t.autofit = False
    ncol = len(rows[0])
    first = 4.5 if rows[0][0] in ("Диод", "Амплитуда генератора, В") else 2.0
    widths = [first] + [(17.0 - first) / (ncol - 1)] * (ncol - 1)
    t._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders xmlns:w="{W_NS}">'
        + "".join(f'<w:{b} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                  for b in ("top", "left", "bottom", "right", "insideH", "insideV"))
        + "</w:tblBorders>"))
    for gc, wcm in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(wcm * 567)))
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Cm(widths[j])
            par = cell.paragraphs[0]
            pf = par.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.left_indent = Cm(0)
            par.alignment = 1
            r = par.add_run(txt)
            if i == 0:
                r.bold = True
            if txt == NONE:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                r.bold = True


def set_runs(par, texts):
    """первые run'ы получают тексты, остальные очищаются"""
    for r, t in zip(par.runs, texts + [""] * len(par.runs)):
        r.text = t


def titul(doc):
    for p in doc.paragraphs:
        if p.text.startswith("ЛАБОРАТОРНАЯ РАБОТА №"):
            set_runs(p, ["ЛАБОРАТОРНАЯ РАБОТА №", "5"])
        elif "СПЕКТРЫ СИГНАЛОВ" in p.text:
            for r in p.runs:
                if r.text == "СПЕКТРЫ СИГНАЛОВ":
                    r.text = "ПОЛУПРОВОДНИКОВЫЕ ДИОДЫ"


def drop_unused_images(doc):
    body_xml = doc.element.xml
    part = doc.part
    for rid, rel in list(part.rels.items()):
        if "image" in rel.reltype and f'"{rid}"' not in body_xml:
            del part.rels[rid]


def main():
    if DOCX.exists():
        shutil.copy(DOCX, LAB / f"lab5_before_sborka_{time.strftime('%H%M%S')}.docx")
    sobrat_vstavlennye()
    shutil.copy(SHABLON, DOCX)
    doc = docx.Document(str(DOCX))
    body = doc.element.body

    start = next(p for p in doc.paragraphs
                 if p.style.name == "ОТЧЁТ_Заголовок_1" and p.text.strip() == "Цель работы")
    el, sect = start._p, body.find(qn("w:sectPr"))
    while el is not None and el is not sect:
        nxt = el.getnext()
        body.remove(el)
        el = nxt
    titul(doc)
    drop_unused_images(doc)

    GRAFIKI.update({"GPR": grafik_pryam(), "GOBR": grafik_obr(), "GST": grafik_stab()})

    # --- 1. Цель
    doc.add_paragraph("Цель работы", style="ОТЧЁТ_Заголовок_1")
    doc.add_paragraph(
        "Цель работы – экспериментально исследовать вольт-амперные характеристики (ВАХ) "
        "диодов; изучить основные свойства, проводимость и параметры диодов; исследовать "
        "практическое применение диодов в источниках питания.", style="ОТЧЁТ_Текст_14")

    # --- 2. Приборы + теория
    doc.add_paragraph("Приборы и оборудование", style="ОТЧЁТ_Заголовок_1")
    doc.add_paragraph("Для снятия ВАХ используется стенд (рисунок 2.1), в его состав входят:")
    for b in [
        "Источник постоянного напряжения V1 (10, 15 или 50 В), подключаемый ключом S1;",
        "Потенциометр R1 (1 кОм) – плавно меняет напряжение, подаваемое на диод;",
        "Мультиметр XMM1 в режиме амперметра постоянного тока – измеряет ток через диод;",
        "Мультиметр XMM2 в режиме вольтметра постоянного тока – измеряет напряжение на диоде;",
        "Переключатели S2, S3 – прямое или обратное включение диода; ключи S4–S6 – выбор диода;",
        "Исследуемые диоды: D1 – 1BH62, D2 – 1N1199C, D3 – стабилитрон 02BZ2.2.",
    ]:
        doc.add_paragraph(b, style="ОТЧЁТ_Список_маркированный")
    add_fig(doc, "S01", "Схема стенда для снятия ВАХ диодов", "S01: рисунок 2.1")
    doc.add_paragraph("Для исследования источника питания собрана схема (рисунок 2.2):")
    for b in [
        "Функциональный генератор XFG1 – источник синусоидального напряжения;",
        "Трансформатор T1;",
        "Диодный мост D1 (3N259) – двухполупериодный выпрямитель;",
        "Конденсатор C1 (300 мкФ) – сглаживающий фильтр;",
        "Резистор R1 (5 кОм) – нагрузка;",
        "Двухканальный осциллограф XSC1: канал A – вход, канал B – напряжение на нагрузке.",
    ]:
        doc.add_paragraph(b, style="ОТЧЁТ_Список_маркированный")
    add_fig(doc, "S02", "Схема источника питания", "S02: рисунок 2.2")

    doc.add_paragraph("Краткие теоретические сведения", style="ОТЧЁТ_Заголовок_2")
    for it in TEORIYA:
        (add_text if it[0] == "p" else add_eq)(doc, it[1])

    # --- 3. Ход работы
    items = content()
    fig_no, tbl_no, nf, nt = {}, {}, 0, 0
    for it in items:
        if it[0] == "fig":
            nf += 1
            fig_no[it[1]] = f"3.{nf}"
        elif it[0] == "tbl":
            nt += 1
            tbl_no[it[1]] = f"3.{nt}"
    refs = {**fig_no, **tbl_no}

    doc.add_paragraph("Ход работы", style="ОТЧЁТ_Заголовок_1")
    for it in items:
        kind = it[0]
        if kind == "h2":
            doc.add_paragraph(it[1], style="ОТЧЁТ_Заголовок_2")
        elif kind == "p":
            add_text(doc, it[1].format(**refs))
        elif kind == "eq":
            add_eq(doc, it[1])
        elif kind == "fig":
            _, key, src, cap = it
            add_fig(doc, src, cap, f"{src}: рисунок {fig_no[key]}")
        elif kind == "tbl":
            add_table(doc, it[2], it[3])

    # --- 4. Выводы
    doc.add_paragraph("Выводы", style="ОТЧЁТ_Заголовок_1")
    for i, v in enumerate(vyvody(), 1):
        add_text(doc, f"{i}. {v}")

    doc.save(str(DOCX))
    print("рисунков в ходе работы:", nf, "таблиц:", nt,
          "вставлено руками:", list(VSTAVLENO))
    for key, no in fig_no.items():
        print(f"  {no}  {key}")


if __name__ == "__main__":
    main()
