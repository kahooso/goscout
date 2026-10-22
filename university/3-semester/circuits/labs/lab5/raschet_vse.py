"""ЛР-5: все таблицы и графики — РАСЧЁТ по моделям из базы Multisim (решение пользователя
2026-10-14 не мерить вручную). Реально сняты только точки 50 % прямых ветвей (zamery.py).

Модели (MSCOMP_S.PRD):
  1BH62    D  IS=5.95e-6  N=4.031  RS=0.02677  BV=120
  1N1199C  D  IS=9.94e-6  N=2      RS=0.00643  BV=100
  02BZ2.2  ZENER v_breakdown=2.2 i_breakdown=10e-3 r_breakdown=15 i_rev=5e-9 i_sat=1e-12 n_forward=1.7874
IS каждой модели чуть подогнан, чтобы она прошла точно через замер при 50 % (поправка < 1 %).

Схема стенда: потенциометр на x — источник E·x с внутренним сопротивлением 1000·x·(1−x) Ом.
Источник питания: генератор Amplitude A → на первичной обмотке 2A (между «+» и «−»),
трансформатор 1:1, мост — два открытых диода по UD, C1 ‖ 5 кОм.

Запуск: python raschet_vse.py → «Таблицы ЛР-5 (v3).docx» + grafik_*.png + числа для текста.
"""
import math
import sys
from pathlib import Path

import docx
import matplotlib
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

LAB = Path(__file__).parent
sys.path.insert(0, str(LAB))
import zamery as Z  # noqa: E402

VT = 0.025865          # 27 °C
IBV = 1e-3             # SPICE по умолчанию
NAMES = {"D1": "1BH62", "D2": "1N1199C", "D3": "02BZ2.2"}
MODEL = {              # n, Is, Rs, BV
    "D1": (4.031, 5.95e-6, 0.02677, 120.0),
    "D2": (2.0, 9.94e-6, 0.00643, 100.0),
    "D3": (1.7874, 1e-12, 0.0, None),
}
VZ, IZ, RZ, I_REV = 2.2, 10e-3, 15.0, 5e-9


# --- модели ------------------------------------------------------------------
def kalibr():
    """подогнать Is под замер 50 %"""
    m = {}
    for d, (n, i_s, rs, bv) in MODEL.items():
        _, u, i = Z.PRYAM[d][Z.POZ_PRYAM.index(50)]
        m[d] = (n, i / (math.exp((u - i * rs) / (n * VT)) - 1), rs, bv)
    return m


M = kalibr()


def i_pryam(d, u):
    n, i_s, rs, _ = M[d]
    # u = Uп + I·Rs → решаем по Uп
    lo, hi = 0.0, u
    for _ in range(100):
        up = (lo + hi) / 2
        i = i_s * (math.exp(up / (n * VT)) - 1)
        if up + i * rs > u:
            hi = up
        else:
            lo = up
    return i_s * (math.exp(lo / (n * VT)) - 1)


def i_obr(d, u):
    """обратный ток (u > 0 — обратное напряжение)"""
    if d == "D3":
        return I_REV * (1 - math.exp(-u / VT)) + IZ * math.exp((u - VZ) / (IZ * RZ))
    n, i_s, rs, bv = M[d]
    i = K_OBR * i_s * (1 - math.exp(-u / (n * VT)))
    if u > bv - 1:
        i += IBV * math.exp((u - bv) / VT)
    return i


def tochka(func, d, e, x, rpot=1000):
    """ползунок x (доля от НИЖНЕГО вывода), источник e: решаем U + I(U)·Rт = E·x"""
    et, rt = e * x, rpot * x * (1 - x)
    lo, hi = 0.0, et
    for _ in range(200):
        u = (lo + hi) / 2
        if u + func(d, u) * rt > et:
            hi = u
        else:
            lo = u
    u = (lo + hi) / 2
    return u, func(d, u)


# --- форматирование ----------------------------------------------------------
def f(x, nd=3):
    return f"{x:.{nd}f}".replace(".", ",")


def g(x, sig=3):
    if x == 0:
        return "0"
    s = f"{x:#.{sig}g}"
    if "e" in s:
        e = int(s.split("e")[1])
        s = f"{x:.{sig - 1 - e}f}" if e < 0 else f"{x:.0f}"
    return s.rstrip(".").replace(".", ",")


# --- расчёт ------------------------------------------------------------------
def pryamye():
    res = {}
    for d in NAMES:
        rows = []
        for p in Z.POZ_PRYAM:
            zam = Z.PRYAM[d][Z.POZ_PRYAM.index(p)]
            if zam[1] is not None:
                rows.append((p, zam[1], zam[2]))
            else:
                rows.append((p, *tochka(i_pryam, d, 10, p / 100)))
        res[d] = rows
    return res


V_OBR = 150
R_POT_OBR = 100e3      # при 150 В потенциометр 1 кОм сгорает (22,5 Вт) — поставлен 100 кОм
POZ_OBR = [2, 10, 20, 30, 40, 50, 60, 70, 80, 90]
# Multisim отсчитывает % от ВЕРХНЕГО вывода: подпись у R1 = 100 − x.
# Замер D1 при подписи 90 %: 14,932 В / 7,105 мкА, а модель даёт Is = 5,95 мкА →
# обратный ток модели умножаем на поправку замера.
K_OBR = 7.105e-6 / 5.95e-6
ZAMER_OBR = {("D1", 10): (14.932, 7.105e-6), ("D1", 90): (120.429, 1.62e-3),
             ("D2", 90): (100.094, 3.88e-3)}


def obratnye():
    res = {}
    for d in ("D1", "D2"):
        rows = []
        for p in POZ_OBR:
            if (d, p) in ZAMER_OBR:
                rows.append((p, *ZAMER_OBR[(d, p)]))
            else:
                rows.append((p, *tochka(i_obr, d, V_OBR, p / 100, R_POT_OBR)))
        res[d] = rows
    return res


POZ_ST = [2, 5, 10, 12, 14, 15, 16, 17, 18, 20, 25, 30, 40, 60]


def stab():
    return [(p, *tochka(i_obr, "D3", 15, p / 100)) for p in POZ_ST]


def sopr(rows, p):
    k = [r[0] for r in rows].index(p)
    a, c, b = rows[k - 1], rows[k], rows[k + 1]
    return c[1], c[2], c[1] / c[2], (b[1] - a[1]) / (b[2] - a[2])


def u_pri_toke(rows, i_t):
    pts = sorted(rows, key=lambda r: r[2])
    for a, b in zip(pts, pts[1:]):
        if a[2] <= i_t <= b[2]:
            return a[1] + (b[1] - a[1]) * (i_t - a[2]) / (b[2] - a[2])


# источник питания
UD = 0.65          # падение на одном диоде моста при токах единицы мА
RN, RD = 5000.0, 1000.0


def u_stab(uc):
    """напряжение на стабилитроне при Uc на конденсаторе (Rд = 1 кОм, нагрузка 5 кОм)"""
    lo, hi = 0.0, uc
    for _ in range(100):
        u = (lo + hi) / 2
        if (uc - u) / RD > i_obr("D3", u) + u / RN:
            lo = u
        else:
            hi = u
    return u


def istochnik(amp, c, fr, so_stab):
    """(Uпик на C, размах пульсаций на C, среднее на выходе, размах на выходе)"""
    upk = 2 * amp - 2 * UD
    if c == 0:
        return upk, upk, 2 / math.pi * upk, upk
    if so_stab:
        uz = u_stab(upk)
        i_load = (upk - uz) / RD
    else:
        i_load = upk / RN
    rip = i_load / (2 * fr * c)
    uc = upk - rip / 2
    if not so_stab:
        return upk, rip, uc, rip
    uz = u_stab(uc)
    rdyn = IZ * RZ / i_obr("D3", uz)          # dU/dI экспоненты
    k = (rdyn * RN / (rdyn + RN)) / (RD + rdyn * RN / (rdyn + RN))
    return upk, rip, uz, rip * k


# --- вывод -------------------------------------------------------------------
def main():
    PR, OB, ST = pryamye(), obratnye(), stab()

    d = docx.Document()
    sty = d.styles["Normal"]
    sty.font.name = "Times New Roman"
    sty.font.size = Pt(14)
    sty.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    def table(cap, rows, widths):
        p = d.add_paragraph(cap)
        p.paragraph_format.keep_with_next = True
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, txt in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Cm(widths[j])
                par = cell.paragraphs[0]
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run(txt).bold = i == 0
        d.add_paragraph()

    # 3.1
    rows = [["R1, %", "U1, В", "I1, мА", "U2, В", "I2, мА", "U3, В", "I3, мА"]]
    for k, p in enumerate(Z.POZ_PRYAM):
        row = [str(100 - p)]
        for dd in NAMES:
            _, u, i = PR[dd][k]
            row += [f(u), g(i * 1e3)]
        rows.append(row)
    table("Таблица 3.1 – Прямые ветви ВАХ диодов D1, D2, D3", rows, [2.0] + [2.4] * 6)

    # 3.2
    rows = [["Диод", "U, В", "I, мА", "Rст = U/I, Ом", "rдиф = ΔU/ΔI, Ом"]]
    SOPR = {}
    for dd in NAMES:
        u, i, rst, rd = sopr(PR[dd], 50)
        SOPR[dd] = (u, i, rst, rd)
        rows.append([f"{dd} ({NAMES[dd]}), прямое", f(u), g(i * 1e3), g(rst), g(rd)])
    u, i, rst, rd = sopr(ST, 30)
    SOPR["D3пр"] = (u, i, rst, rd)
    rows.append(["D3 (02BZ2.2), пробой", f(u), g(i * 1e3), g(rst), g(rd)])
    table("Таблица 3.2 – Статические и динамические сопротивления в рабочих точках",
          rows, [5.0, 2.3, 2.3, 3.4, 4.0])

    # 3.3
    rows = [["R1, %", "U1, В", "I1, мА", "U2, В", "I2, мА"]]
    for k, p in enumerate(POZ_OBR):
        rows.append([str(100 - p), f(OB["D1"][k][1], 2), g(OB["D1"][k][2] * 1e3),
                     f(OB["D2"][k][1], 2), g(OB["D2"][k][2] * 1e3)])
    table(f"Таблица 3.3 – Обратные ветви ВАХ диодов D1 и D2 (V1 = {V_OBR} В)", rows,
          [2.0, 3.5, 3.5, 3.5, 3.5])

    # 3.4
    rows = [["R1, %", "U, В", "I, мА"]]
    for p, u, i in ST:
        rows.append([str(100 - p), f(u), g(i * 1e3)])
    table("Таблица 3.4 – Обратная ветвь ВАХ стабилитрона D3 (V1 = 15 В)", rows, [3.0, 5.0, 5.0])

    # 3.5
    rows = [["Амплитуда генератора, В", "Uвых без стабилитрона, В", "Uвых со стабилитроном, В"]]
    T5 = {}
    for a in (3, 4, 5, 6, 7):
        bez = istochnik(a, 300e-6, 50, False)[2]
        so = istochnik(a, 300e-6, 50, True)[2]
        T5[a] = (bez, so)
        rows.append([str(a), f(bez, 2), f(so, 3)])
    table("Таблица 3.5 – Выходное напряжение источника при разной амплитуде входного сигнала",
          rows, [5.0, 5.5, 5.5])
    out = LAB / "Таблицы ЛР-5 (v3).docx"
    d.save(out)

    # графики
    def ris(name, series, xl, yl, quad3=False):
        fig, ax = plt.subplots(figsize=(7, 4.5), dpi=150)
        for lab, xs, ys, c in series:
            if quad3:
                xs, ys = [-v for v in xs], [-v for v in ys]
            ax.plot(xs, ys, "o-", ms=4, color=c, label=lab)
        ax.set_xlabel(xl)
        ax.set_ylabel(yl)
        ax.grid(True, alpha=0.4)
        ax.legend()
        fig.tight_layout()
        fig.savefig(LAB / name)
        plt.close(fig)

    C = {"D1": "#1f77b4", "D2": "#d62728", "D3": "#2ca02c"}
    ris("grafik_1_pryamye.png",
        [(f"{dd} ({NAMES[dd]})", [r[1] for r in PR[dd]], [r[2] * 1e3 for r in PR[dd]], C[dd])
         for dd in NAMES], "U, В", "I, мА")
    ris("grafik_2_obratnye.png",
        [(f"{dd} ({NAMES[dd]})", [r[1] for r in OB[dd]], [r[2] * 1e3 for r in OB[dd]], C[dd])
         for dd in OB], "U, В", "I, мА", quad3=True)
    ris("grafik_3_stabilitron.png",
        [("D3 (02BZ2.2), обратное включение", [r[1] for r in ST], [r[2] * 1e3 for r in ST],
          C["D3"])], "U, В", "I, мА", quad3=True)

    # числа для текста
    print("Файл:", out.name)
    print("\n[3.1] U при 1 мА:", {dd: f(u_pri_toke(PR[dd], 1e-3), 2) for dd in NAMES})
    print("[3.2] Rст / rдиф:", {k: (g(v[2]), g(v[3])) for k, v in SOPR.items()})
    for dd in OB:
        p, u, i = OB[dd][POZ_OBR.index(60)]
        print(f"[3.3] {dd}: до пробоя (подпись 40 %) U={f(u, 1)} В I={g(i * 1e6)} мкА "
              f"Rобр={g(u / i / 1e6)} МОм ≈ {g(u / i / SOPR[dd][2])} раз больше прямого")
    print("[3.4] Uст при 10 мА:", f(u_pri_toke(ST, 10e-3), 3),
          "| пробой от", g(ST[POZ_ST.index(15)][2] * 1e3), "до", g(ST[-1][2] * 1e3),
          "мА: U", f(ST[POZ_ST.index(15)][1]), "→", f(ST[-1][1]))
    for lab, c, fr, s in (("без C", 0, 50, False), ("300 мкФ 50 Гц", 300e-6, 50, False),
                          ("30 мкФ 50 Гц", 30e-6, 50, False), ("30 мкФ 500 Гц", 30e-6, 500, False),
                          ("стаб 300 мкФ 50 Гц", 300e-6, 50, True)):
        upk, rip, uout, ripout = istochnik(5, c, fr, s)
        print(f"[4] {lab}: Uпик={f(upk, 2)} В, пульсации на C={f(rip, 3)} В, "
              f"Uвых={f(uout, 3)} В, пульсации на выходе={f(ripout, 4)} В")
    print("[3.5]", {a: (f(b, 2), f(s, 3)) for a, (b, s) in T5.items()})


if __name__ == "__main__":
    main()
