# -*- coding: utf-8 -*-
"""ЛР-3 «Электрические фильтры». Расчёт всех таблиц АЧХ и выгрузка в Word."""
import math, io, sys, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# ---------- модель ----------
def ZC(f, C): return 1/(1j*2*math.pi*f*C)
def ZL(f, L): return 1j*2*math.pi*f*L
def par(a, b): return (a*b)/(a+b)
R = 1000.0

def rc_hp(f, C, U=10): return abs(U*R/(R+ZC(f, C)))
def rc_lp(f, C, U=10): z = ZC(f, C); return abs(U*z/(R+z))
def rl_hp(f, L, U=10): z = ZL(f, L); return abs(U*z/(R+z))
def rl_lp(f, L, U=10): return abs(U*R/(R+ZL(f, L)))

def t_hp(f, C, U=10):                      # V-R1-C1-[R3]-C2-[R4]
    z1 = z2 = ZC(f, C)
    zb = par(R, z2+R)
    return abs(U*zb/(R+z1+zb) * R/(z2+R))

def p_lp(f, C, U=10):                      # V-R1-[C1]-R3-[C2||R4]
    z1 = z2 = ZC(f, C)
    zb = par(z2, R)
    za = par(z1, R+zb)
    return abs(U*za/(R+za) * zb/(R+zb))

def t_lp(f, L, U=10):                      # V-R1-L1-[R3]-L2-[R4]  (зеркало t_hp)
    z1 = z2 = ZL(f, L)
    zb = par(R, z2+R)
    return abs(U*zb/(R+z1+zb) * R/(z2+R))

def p_hp(f, L, U=10):                      # V-R1-[L1]-R3-[L2||R4] (зеркало p_lp)
    z1 = z2 = ZL(f, L)
    zb = par(z2, R)
    za = par(z1, R+zb)
    return abs(U*za/(R+za) * zb/(R+zb))

def notch(f, C, L=405e-6, U=10):           # V-R4-[L||C]-R1
    return abs(U*R/(R+par(ZL(f, L), ZC(f, C))+R))

def bp(f, C, L=9e-3, U=0.5):               # V-R2-[(L||C)||R1]
    zs = par(par(ZL(f, L), ZC(f, C)), R)
    return abs(U*zs/(R+zs))

# ---------- сетки частот ----------
G1  = [200, 500, 1000, 2000, 5000, 8000, 16000, 30000, 60000, 150000]
N10 = [20000, 50000, 65000, 74000, 79000, 84000, 95000, 120000, 150000, 200000]
N20 = [15000, 35000, 46000, 52000, 56000, 60000, 68000, 85000, 120000, 200000]
B10 = [2000, 5000, 8000, 12000, 17000, 22000, 30000, 50000, 80000, 150000]
B20 = [2000, 5000, 8000, 12000, 17000, 22000, 30000, 50000, 80000, 150000]

def build(fn, p, grid, U):
    rows = [[f, fn(f, p), fn(f, p)/U] for f in grid]
    kmax = max(r[2] for r in rows)
    for r in rows: r.append(r[2]/kmax)
    return rows, U

SECTIONS = [
 ("1. Исследование простейших RC- и RL-фильтров высоких частот", [
   ("1.1", "RC-ФВЧ при C1 = 10 нФ, R1 = 1 кОм",  build(rc_hp, 10e-9, G1, 10), "Граничная частота fгр = 1/(2πRC) = 15,9 кГц."),
   ("1.2", "RC-ФВЧ при C2 = 20 нФ, R1 = 1 кОм",  build(rc_hp, 20e-9, G1, 10), "Граничная частота fгр = 1/(2πRC) = 7,96 кГц."),
   ("1.3", "RL-ФВЧ при L1 = 10 мГн, R1 = 1 кОм", build(rl_hp, 10e-3, G1, 10), "Граничная частота fгр = R/(2πL) = 15,9 кГц."),
   ("1.4", "RL-ФВЧ при L2 = 20 мГн, R1 = 1 кОм", build(rl_hp, 20e-3, G1, 10), "Граничная частота fгр = R/(2πL) = 7,96 кГц."),
 ]),
 ("2. Исследование простейших RC- и RL-фильтров низких частот", [
   ("2.1", "RC-ФНЧ при C1 = 10 нФ, R1 = 1 кОм",  build(rc_lp, 10e-9, G1, 10), "Граничная частота fгр = 1/(2πRC) = 15,9 кГц."),
   ("2.2", "RC-ФНЧ при C2 = 20 нФ, R1 = 1 кОм",  build(rc_lp, 20e-9, G1, 10), "Граничная частота fгр = 1/(2πRC) = 7,96 кГц."),
   ("2.3", "RL-ФНЧ при L1 = 10 мГн, R1 = 1 кОм", build(rl_lp, 10e-3, G1, 10), "Граничная частота fгр = R/(2πL) = 15,9 кГц."),
   ("2.4", "RL-ФНЧ при L2 = 20 мГн, R1 = 1 кОм", build(rl_lp, 20e-3, G1, 10), "Граничная частота fгр = R/(2πL) = 7,96 кГц."),
 ]),
 ("3. Исследование сложных Т- и П-образных ФВЧ и ФНЧ", [
   ("3.1", "Т-образный ФВЧ при C1 = C2 = 10 нФ, R1 = R2 = R3 = 1 кОм", build(t_hp, 10e-9, G1, 10),
    "Граничная частота по уровню 0,707 составляет 17,3 кГц. В полосе пропускания Kmax = 0,331: часть сигнала теряется на резисторах схемы."),
   ("3.2", "Т-образный ФВЧ при C1 = C2 = 20 нФ, R1 = R2 = R3 = 1 кОм", build(t_hp, 20e-9, G1, 10),
    "Граничная частота по уровню 0,707 составляет 8,7 кГц. Kmax = 0,333."),
   ("3.3", "Т-образный ФНЧ при L1 = L2 = 10 мГн, R1 = R2 = R3 = 1 кОм", build(t_lp, 10e-3, G1, 10),
    "Граничная частота по уровню 0,707 составляет 14,5 кГц. Kmax = 0,333."),
   ("3.4", "Т-образный ФНЧ при L1 = L2 = 20 мГн, R1 = R2 = R3 = 1 кОм", build(t_lp, 20e-3, G1, 10),
    "Граничная частота по уровню 0,707 составляет 7,3 кГц. Kmax = 0,333."),
   ("3.5", "П-образный ФВЧ при L1 = L2 = 10 мГн, R1 = R2 = R3 = 1 кОм", build(p_hp, 10e-3, G1, 10),
    "Граничная частота по уровню 0,707 составляет 17,3 кГц. Kmax = 0,331."),
   ("3.6", "П-образный ФВЧ при L1 = L2 = 20 мГн, R1 = R2 = R3 = 1 кОм", build(p_hp, 20e-3, G1, 10),
    "Граничная частота по уровню 0,707 составляет 8,7 кГц. Kmax = 0,333."),
   ("3.7", "П-образный ФНЧ при C1 = C2 = 5 нФ, R1 = R2 = R3 = 1 кОм",  build(p_lp, 5e-9, G1, 10),
    "Граничная частота по уровню 0,707 составляет 29,0 кГц. Kmax = 0,333."),
   ("3.8", "П-образный ФНЧ при C1 = C2 = 10 нФ, R1 = R2 = R3 = 1 кОм", build(p_lp, 10e-9, G1, 10),
    "Граничная частота по уровню 0,707 составляет 14,5 кГц. Kmax = 0,333."),
 ]),
 ("4. Исследование простейших резонансных фильтров", [
   ("4.1", "Режекторный фильтр при L1 = 405 мкГн, C1 = 10 нФ",  build(notch, 10e-9, N10, 10),
    "Резонансная частота f0 = 1/(2π√(LC)) = 79,1 кГц. На этой частоте параллельный контур имеет максимальное сопротивление и сигнал на выход не проходит."),
   ("4.2", "Режекторный фильтр при L1 = 405 мкГн, C1 = 20 нФ",  build(notch, 20e-9, N20, 10),
    "Резонансная частота f0 = 55,9 кГц. Увеличение ёмкости в два раза смещает режекцию вниз по частоте в √2 раз."),
   ("4.3", "Полосовой фильтр при L1 = 9 мГн, C1 = 10 нФ",       build(bp, 10e-9, B10, 0.5),
    "Резонансная частота f0 = 16,8 кГц. Напряжение источника Uвх = 0,5 В. Kmax = 0,50."),
   ("4.4", "Полосовой фильтр при L1 = 9 мГн, C1 = 20 нФ",       build(bp, 20e-9, B20, 0.5),
    "Резонансная частота f0 = 11,9 кГц. Напряжение источника Uвх = 0,5 В. Kmax = 0,50."),
 ]),
]

# ---------- формат чисел ----------
def fF(f):  return f"{f:,}".replace(",", " ")
def fU(v):  return f"{v*1000:.1f}".replace(".", ",") + " мВ" if v < 1 else f"{v:.3f}".replace(".", ",") + " В"
def fUin(u): return (f"{u:.3f}" if u >= 1 else f"{u:.3f}").replace(".", ",")
def fK(v):  return f"{v:.4f}".replace(".", ",")

# ---------- Word ----------
doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)

def para(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False, after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'; r.font.size = Pt(size)
    r.bold = bold; r.italic = italic
    return p

para("Лабораторная работа № 3", 14, True, WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("ЭЛЕКТРИЧЕСКИЕ ФИЛЬТРЫ", 14, True, WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("Таблицы результатов", 12, False, WD_ALIGN_PARAGRAPH.CENTER, after=14)

para("Обозначения в таблицах:", 12, True)
for s in ("F — частота гармонического сигнала на входе фильтра, Гц;",
          "Uвх — напряжение на входе фильтра (напряжение источника), В;",
          "Uвых — напряжение на выходе фильтра (на нагрузочном резисторе);",
          "K = Uвых/Uвх — коэффициент передачи фильтра;",
          "K̂ = K/Kmax — коэффициент передачи, приведённый к максимальному значению."):
    para("— " + s, 12, after=0)
para("Все характеристики приведены к максимальному значению АЧХ, что позволяет "
     "сравнивать фильтры между собой. Граничной считается частота, на которой K̂ = 0,707.", 12, after=14)

HDR = ["F, Гц", "Uвх, В", "Uвых", "K = Uвых/Uвх", "K̂ = K/Kmax"]

for sec_title, tables in SECTIONS:
    para(sec_title, 13, True, after=10)
    for num, title, (rows, U), note in tables:
        para(f"Таблица {num} — {title}", 12, after=4)
        t = doc.add_table(rows=1, cols=5)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(HDR):
            c = t.rows[0].cells[i]
            c.text = ""
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h); r.bold = True; r.font.size = Pt(11); r.font.name = 'Times New Roman'
        for f, u, k, kn in rows:
            cells = t.add_row().cells
            for i, v in enumerate([fF(f), fUin(U), fU(u), fK(k), fK(kn)]):
                cells[i].text = ""
                p = cells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(v); r.font.size = Pt(11); r.font.name = 'Times New Roman'
        para(note, 11, italic=True, after=14)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Таблицы ЛР-3.docx")
doc.save(out)
print("сохранено:", out)
print("таблиц:", sum(len(t) for _, t in SECTIONS))
